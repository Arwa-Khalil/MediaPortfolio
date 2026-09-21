from docx import Document

doc = Document('media-portfolio-backend/API-Contract-v1.0.docx')

# Let's find the table under 4.2 GET /admin/videos
# We will iterate through paragraphs to find '4.2 GET /admin/videos', then find the next table.
found_42 = False
table_to_edit = None

for i, element in enumerate(doc.element.body):
    if element.tag.endswith('p'):
        p = [p for p in doc.paragraphs if p._p == element][0]
        if '4.2 GET /admin/videos' in p.text:
            found_42 = True
    elif element.tag.endswith('tbl') and found_42:
        table_to_edit = [t for t in doc.tables if t._tbl == element][0]
        break

if table_to_edit:
    # Add new rows
    row1 = table_to_edit.add_row()
    row1.cells[0].text = 'cloudflareThumbnailImageId'
    row1.cells[1].text = 'string (nullable)'
    row1.cells[2].text = 'Cloudflare image ID for the custom thumbnail'

    row2 = table_to_edit.add_row()
    row2.cells[0].text = 'thumbnailUrl'
    row2.cells[1].text = 'string (nullable)'
    row2.cells[2].text = 'Absolute URL to the thumbnail image (custom or auto-generated)'

    doc.save('media-portfolio-backend/API-Contract-v1.0.docx')
    print("Document updated successfully.")
else:
    print("Could not find table for 4.2 GET /admin/videos")
