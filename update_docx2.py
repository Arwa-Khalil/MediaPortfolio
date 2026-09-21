from docx import Document

doc = Document('media-portfolio-backend/API-Contract-v1.0.docx')

# Paragraph index for GET /admin/videos is 63
start_p = doc.paragraphs[63]
end_p = doc.paragraphs[66]

found_start = False
table_to_edit = None

for element in doc.element.body:
    if element.tag.endswith('p'):
        p = [p for p in doc.paragraphs if p._p == element][0]
        if p == start_p:
            found_start = True
        elif p == end_p:
            break
    elif element.tag.endswith('tbl') and found_start:
        table_to_edit = [t for t in doc.tables if t._tbl == element][0]
        break

if table_to_edit:
    # Check if cloudflareThumbnailImageId already exists
    already_has_field = False
    for row in table_to_edit.rows:
        if 'cloudflareThumbnailImageId' in row.cells[0].text:
            already_has_field = True
            break
            
    if not already_has_field:
        row1 = table_to_edit.add_row()
        row1.cells[0].text = 'cloudflareThumbnailImageId'
        row1.cells[1].text = 'string (nullable)'
        row1.cells[2].text = 'Cloudflare image ID for the custom thumbnail'

        row2 = table_to_edit.add_row()
        row2.cells[0].text = 'thumbnailUrl'
        row2.cells[1].text = 'string (nullable)'
        row2.cells[2].text = 'Absolute URL to the thumbnail image (custom or auto-generated)'

        doc.save('media-portfolio-backend/API-Contract-v1.0.docx')
        print("Document updated successfully.")
    else:
        print("Document already has cloudflareThumbnailImageId in the table.")
else:
    print("Could not find table for GET /admin/videos")
